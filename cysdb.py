"""CysDB ligandable-cysteine layer: find the covalent best-in-class PRIZE the patent radar can't —
clinically-validated drug targets that have a ligandable cysteine but NO covalent program yet.

Intersect three signals:
  LIGANDABLE  (CysDB Fig4A-C 'CysDB Ligandable'=yes + #ligandable cysteines from Ligandable Dataset)
              -> a covalent small molecule is chemically possible.
  VALIDATED   (CysDB FDA/ChEMBL/DrugBank drug-target flags) -> the biology is derisked.
  COVALENT-OPEN (gene NOT in our SureChEMBL covalent-patent attribution set) -> window not yet taken.
The sweet spot = LIGANDABLE & FDA-validated & covalent-OPEN.  Writes cysdb_opportunities.json.
"""
import csv
import json
from collections import Counter
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
XLSX = HERE / "NIHMS1893018-supplement-2.xlsx"
_LIGAND_CACHE = HERE / ".cysdb_ligand.json"
_HGNC_URL = "https://storage.googleapis.com/public-download-files/hgnc/tsv/tsv/hgnc_complete_set.txt"


def ligandability():
    """sym -> {'ligandable': bool, 'n_lig_cys': int} from RAW CysDB (cached) — for ANY protein, not just the
    filtered opportunities. Lets the dossier ask 'is target X covalently ligandable per chemoproteomics?'."""
    import json as _json
    if _LIGAND_CACHE.exists():
        return _json.load(open(_LIGAND_CACHE))
    import urllib.request
    if not Path("/tmp/hgnc.txt").exists():                       # complete set has the uniprot_ids column
        urllib.request.urlretrieve(_HGNC_URL, "/tmp/hgnc.txt")
    uni2sym = {}
    for r in csv.DictReader(open("/tmp/hgnc.txt"), delimiter="\t"):
        if r.get("locus_group") == "protein-coding gene":
            for u in (r.get("uniprot_ids") or "").split("|"):
                if u.strip():
                    uni2sym[u.strip()] = r["symbol"]
    wb = openpyxl.load_workbook(str(XLSX), read_only=True)
    ligandable = {}
    it = wb["Fig4A-C"].iter_rows(values_only=True); next(it)
    for pid, fda, chembl, drugbank, lig in it:
        if pid:
            ligandable[pid] = lig == "yes"
    ncys = Counter()
    it = wb["Ligandable Dataset"].iter_rows(min_col=1, max_col=4, values_only=True); next(it)
    for pid, cysid, resid, lig in it:
        if pid and lig == "yes":
            ncys[pid] += 1
    out = {}
    for pid, lig in ligandable.items():
        sym = uni2sym.get(pid)
        if sym:
            b = out.setdefault(sym, {"ligandable": False, "n_lig_cys": 0})
            b["ligandable"] = b["ligandable"] or lig
            b["n_lig_cys"] = max(b["n_lig_cys"], ncys.get(pid, 0))
    _json.dump(out, open(_LIGAND_CACHE, "w"))
    return out


def run():
    # UniProt -> HGNC symbol (protein-coding only)
    uni2sym = {}
    for r in csv.DictReader(open("/tmp/hgnc.txt"), delimiter="\t"):
        if r.get("locus_group") == "protein-coding gene":
            for u in (r.get("uniprot_ids") or "").split("|"):
                if u.strip():
                    uni2sym[u.strip()] = r["symbol"]

    wb = openpyxl.load_workbook(str(XLSX), read_only=True)
    # Fig4A-C: per-protein drug-target + ligandable flags
    prot = {}
    it = wb["Fig4A-C"].iter_rows(values_only=True); next(it)
    for pid, fda, chembl, drugbank, lig in it:
        if pid:
            prot[pid] = {"fda": fda == "yes", "chembl": chembl == "yes",
                         "drugbank": drugbank == "yes", "ligandable": lig == "yes"}
    # Ligandable Dataset: count ligandable cysteines per protein (cols: proteinid, cysteineid, resid, ligandable)
    ncys = Counter()
    it = wb["Ligandable Dataset"].iter_rows(min_col=1, max_col=4, values_only=True); next(it)
    for pid, cysid, resid, lig in it:
        if pid and lig == "yes":
            ncys[pid] += 1

    # genes with a covalent patent — only a TITLE/CLAIMS attribution counts as "covalent-closed"; a lone
    # body-mention is too weak (the RAF1 false-positive: flagged closed by one quinazoline patent's body
    # pathway-mention of RAF1, when no covalent RAF program actually exists).
    from attribute import trusted                          # post-audit grade: drops combination/viral/material/degrader noise
    tg = json.load(open(HERE / "patent_targets.json"))
    cov_patented = {d["sym"] for d in tg.values() if d.get("sym") and trusted(d)}
    cov_body_only = {d["sym"] for d in tg.values() if d.get("sym")} - cov_patented
    # pan-essential genes (Hart CEG2 + DepMap common essentials): covalently drugging a housekeeping
    # enzyme (tubulin, TYMS, polymerases, FASN...) = a toxin, not a best-in-class drug -> exclude.
    essential = {ln.split("\t")[0] for ln in open("/tmp/ceg2.txt").read().splitlines()[1:] if ln.strip()}
    essential |= {ln.split(" (")[0].strip() for ln in open("/tmp/depmap_common_ess.csv").read().splitlines()[1:] if ln.strip()}

    rows = []
    for pid, p in prot.items():
        if not p["ligandable"]:
            continue
        sym = uni2sym.get(pid)
        if not sym:
            continue
        rows.append({"uniprot": pid, "sym": sym, "n_lig_cys": ncys.get(pid, 0),
                     "fda": p["fda"], "chembl": p["chembl"], "drugbank": p["drugbank"],
                     "covalent_patented": sym in cov_patented, "core_essential": sym in essential})

    # dedupe by symbol (multiple UniProt isoforms) -> keep max ligandable-cys, OR of flags
    bysym = {}
    for r in rows:
        b = bysym.setdefault(r["sym"], dict(r))
        b["n_lig_cys"] = max(b["n_lig_cys"], r["n_lig_cys"])
        for k in ("fda", "chembl", "drugbank", "covalent_patented", "core_essential"):
            b[k] = b[k] or r[k]
    rows = list(bysym.values())

    val = [r for r in rows if r["fda"] and not r["covalent_patented"]]
    opps = [r for r in val if not r["core_essential"]]            # primary: selectively-druggable
    essential_bucket = [r for r in val if r["core_essential"]]    # secondary: essential (toxic window)
    opps.sort(key=lambda r: -r["n_lig_cys"])
    essential_bucket.sort(key=lambda r: -r["n_lig_cys"])
    json.dump({"opportunities": opps, "essential_bucket": essential_bucket},
              open(HERE / "cysdb_opportunities.json", "w"), indent=1)

    print(f"CysDB ligandable proteins mapped to genes: {len(rows)}")
    print(f"  FDA-validated & ligandable & no covalent patents: {len(val)}")
    print(f"    - core-essential (toxic; separate bucket):      {len(essential_bucket)}")
    recovered = sorted(r["sym"] for r in opps if r["sym"] in cov_body_only)
    print(f"  OPPORTUNITY = + selectively-druggable (non-essential): {len(opps)}")
    print(f"  ...re-surfaced by the title/claims fix (were body-only false-closed): {len(recovered)} {recovered}")
    print("\ntop opportunities by # ligandable cysteines (selectively-druggable, validated, covalent-open):")
    for r in opps[:34]:
        tags = "".join(t for t, f in [("F", r["fda"]), ("C", r["chembl"]), ("D", r["drugbank"])] if f)
        print(f"  {r['sym']:11} lig_cys={r['n_lig_cys']:3}  drugtarget=[{tags}]")


if __name__ == "__main__":
    run()
